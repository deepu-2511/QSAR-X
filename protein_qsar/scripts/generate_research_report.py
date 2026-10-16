"""Generate Markdown and editable Word research reports from results/research_evaluation.py."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
REPORTS = ROOT / "reports"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); element = OxmlElement("w:shd"); element.set(qn("w:fill"), color); tc_pr.append(element)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr(); marker = OxmlElement("w:tblHeader"); marker.set(qn("w:val"), "true"); tr_pr.append(marker)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Light Shading Accent 1"
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, value in zip(header.cells, headers):
        cell.text = value; shade(cell, "17365D")
        for run in cell.paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.font.bold = True; run.font.size=Pt(8)
    for row in rows:
        cells=table.add_row().cells
        for cell,value in zip(cells,row):
            cell.text=str(value)
            for para in cell.paragraphs:
                for run in para.runs: run.font.size=Pt(8)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int=1) -> None:
    p=doc.add_heading(text,level=level)
    for run in p.runs: run.font.color.rgb=RGBColor(23,54,93)


def add_figure(doc: Document, image: Path, caption: str) -> None:
    if not image.exists(): return
    doc.add_picture(str(image), width=Inches(6.3))
    p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs: run.italic=True;run.font.size=Pt(8)


def main() -> None:
    REPORTS.mkdir(exist_ok=True)
    metrics=pd.read_csv(RESULTS/"metrics"/"registered_target_scaffold_test_metrics.csv").sort_values("gene")
    audit=pd.read_csv(RESULTS/"tables"/"dataset_quality_audit.csv").sort_values("gene")
    external=pd.read_csv(RESULTS/"validation"/"external_validation_inventory.csv").sort_values("gene")
    manifest=json.loads((RESULTS/"validation"/"reproducibility_manifest.json").read_text())
    summary=json.loads((RESULTS/"metrics"/"evaluation_summary.json").read_text())
    jak_random=pd.read_csv(RESULTS/"validation"/"JAK_random_vs_scaffold_comparison.csv")
    ad=pd.read_csv(RESULTS/"applicability_domain"/"JAK_scaffold_test_similarity_AD.csv")
    today=date.today().isoformat()
    selected=metrics.loc[metrics.gene.eq("JAK")].iloc[0]
    md=[]
    md += ["# Machine Learning Based QSAR Modeling for Kinase Drug Discovery", "", f"Generated: {today}", "", "## Abstract", "", f"This report evaluates 21 registered molecular activity models in QSAR-X using the saved model artifacts and their recorded scaffold-held-out test partitions. The registered panel contains 20 target-specific ChEMBL kinase models and one pooled JAK-family screening model. The recomputed mean internal scaffold-test R2 was {summary['mean_internal_scaffold_R2']:.3f} and mean RMSE was {summary['mean_internal_scaffold_RMSE']:.3f} pIC50. These are regression performance measures, not biological or clinical accuracy. Strict external validation gates passed for 0 of 21 registered models. Therefore the models are appropriate only for research screening and prioritization.", "", "## Methods", "", "Canonical molecular structures were represented with RDKit Morgan fingerprints (radius 2, 2048 bits). Existing candidate models and their saved split assignments were preserved. The independent partition in this report is a scaffold-held-out internal test, defined by the project’s chirality-aware Bemis-Murcko grouping. Recomputed predictions were used for R2, RMSE, MAE, median absolute error, explained variance, Pearson r, Spearman rho, residual diagnostics, and 300-resample bootstrap confidence intervals for selected metrics. No external-validation claim is made unless the strict external gate report passes.", "", "## Registered Model Performance", "", "| Gene | Model | n test | R2 | RMSE pIC50 | MAE pIC50 | Spearman |", "|---|---:|---:|---:|---:|---:|---:|"]
    for _, r in metrics.iterrows(): md.append(f"| {r.gene} | {r.selected_model} | {int(r.n)} | {r.R2:.3f} | {r.RMSE_pIC50:.3f} | {r.MAE_pIC50:.3f} | {r.Spearman_rho:.3f} |")
    md += ["", "## JAK Family Deep Dive", "", f"The pooled JAK model used {int(selected.n)} scaffold-test compounds in the recomputed evaluation (R2 {selected.R2:.3f}, RMSE {selected.RMSE_pIC50:.3f}, MAE {selected.MAE_pIC50:.3f}). It predicts a pooled pChEMBL activity endpoint. It cannot establish JAK1, JAK2, or JAK3 isoform selectivity because the source rows do not contain a per-row isoform label. The current JAK source also overlaps the locally stored JAK2 training chemistry, therefore it is not an independent external test.", "", "### Random versus scaffold diagnostic", "", "| Split | R2 | RMSE pIC50 | MAE pIC50 |", "|---|---:|---:|---:|"]
    for _,r in jak_random.iterrows(): md.append(f"| {r['split']} | {r.R2:.3f} | {r.RMSE_pIC50:.3f} | {r.MAE_pIC50:.3f} |")
    md += ["", f"The JAK test-set applicability-domain similarity threshold is the empirical 5th percentile of nearest development-set similarity ({manifest['JAK_deep_dive']['AD_test_similarity_threshold_5th_percentile']:.3f}); this is a descriptive internal threshold, not a calibrated error guarantee.", "", "## Dataset Quality and Leakage Audit", "", "Every registered clean input was checked for missing pIC50, duplicate canonical structures, RDKit parsing failures, and exact canonical-structure overlap between development and scaffold-test partitions. All split-overlap counts in the audit table are expected to be zero. The source files retain their curation reports; clean datasets are not silently altered by this evaluator.", "", "## External Validation Status", "", "All models have strict external gates marked false or lack a qualifying external set. Some project reports contain exploratory external-source checks, but missing assay relation, provenance, context, sample-size, independence, or performance conditions prevent promotion. Internal scaffold test performance must not be called external validation.", "", "## Reproducibility", "", f"Random seed: {manifest['random_seed']}. Python: {manifest['python']}; RDKit: {manifest['rdkit']}; scikit-learn: {manifest['scikit_learn']}; NumPy: {manifest['numpy']}; pandas: {manifest['pandas']}. The evaluator stores recomputed predictions, residuals, figures, model SHA256 values, and all tables under results/.", "", "## Limitations", "", "- The target-specific models use molecular fingerprints and do not infer protein binding mechanisms.", "- Scaffold-held-out tests remain internal validation.", "- Prediction errors are in pIC50 and are not individual compound confidence intervals.", "- A high pIC50 means higher potency only for comparable assay conditions; it does not predict safety, efficacy, or clinical success.", "- The pooled JAK model should be used for research screening, not isoform-specific or experimentally verified IC50 claims.", "", "## References", "", "1. OECD. Guidance Document on the Validation of (Quantitative) Structure-Activity Relationship [(Q)SAR] Models. OECD Series on Testing and Assessment No. 69, 2014. https://doi.org/10.1787/9789264085442-en", "2. RDKit documentation. The RDKit Book: Morgan and Feature Morgan fingerprints. https://www.rdkit.org/docs/RDKit_Book.html", "3. Rogers D, Hahn M. Extended-connectivity fingerprints. Journal of Chemical Information and Modeling. 2010;50:742-754. https://doi.org/10.1021/ci100050t", ""]
    (REPORTS/"QSAR_Research_Report.md").write_text("\n".join(md),encoding="utf-8")
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    styles=doc.styles; styles["Normal"].font.name="Aptos";styles["Normal"].font.size=Pt(10)
    title=doc.add_paragraph();title.style="Title";title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run("Machine Learning Based QSAR Modeling for Kinase Drug Discovery")
    subtitle=doc.add_paragraph("Model Development Validation Interpretability and Applicability Domain Assessment");subtitle.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Research report generated from current repository artifacts on {today}").alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    add_heading(doc,"Abstract")
    doc.add_paragraph(f"QSAR-X was audited and evaluated without replacing its existing models. The workflow contains 20 target-specific ChEMBL kinase models and one pooled JAK-family candidate. Recomputed saved-split results give a mean internal scaffold-test R2 of {summary['mean_internal_scaffold_R2']:.3f} and RMSE of {summary['mean_internal_scaffold_RMSE']:.3f} pIC50 across 21 models. R2 and RMSE describe regression fit and error; they are not accuracy. No registered model passed every strict external-validation gate, so this work supports research screening only.")
    add_heading(doc,"Keywords",2);doc.add_paragraph("QSAR; kinase inhibitors; pIC50; Morgan fingerprints; scaffold split; model validation; applicability domain; reproducibility")
    add_heading(doc,"1 Introduction")
    doc.add_paragraph("Quantitative structure-activity relationship modeling maps a molecular representation to an experimentally defined response. A reliable model requires a stated endpoint, a defined applicability domain, transparent validation, and traceable outputs. This report evaluates the existing repository rather than creating a separate model or altering its deployed candidate artifacts.")
    add_heading(doc,"2 Project Audit and Dataset Provenance")
    doc.add_paragraph("The repository contains curation scripts, target-specific training scripts, versioned model artifacts, scaffold assignments, cross-validation comparison tables, external-validation reports, a FastAPI service, and a React interface. Each target model is based on curated ChEMBL IC50 records transformed to pIC50 = 9 - log10(IC50 in nM). The JAK-family candidate uses a user-provided pChEMBL file and is explicitly retained as a pooled endpoint.")
    add_table(doc,["Audit item","Result"],[["Registered models",str(len(metrics))],["Target-specific ChEMBL models","20"],["Pooled JAK screening model","1"],["Fingerprint","Morgan radius 2, 2048 bits"],["Split","Scaffold-grouped development / validation / internal test"],["External gates passed","0"]])
    add_heading(doc,"3 Data Quality Control")
    doc.add_paragraph("The evaluator checked all clean input datasets for missing target values, duplicate canonical structures, RDKit parse failures, and exact structure leakage across saved partitions. The detailed audit CSV records each target and model checksum.")
    add_table(doc,["Gene","Clean structures","Missing pIC50","Duplicate structures","Dev-test overlap"],[[r.gene,str(int(r.n_clean_unique_structures)),str(int(r.missing_pIC50)),str(int(r.duplicate_canonical_smiles)),str(int(r.development_test_structure_overlap))] for _,r in audit.iterrows()])
    add_heading(doc,"4 Molecular Representation and Training")
    doc.add_paragraph("The preserved production candidates use RDKit Morgan fingerprints. The target trainers compare a median baseline, Extra Trees, Random Forest, and SVR using scaffold-grouped cross-validation on the training partition; final candidate selection uses training-CV RMSE and the scaffold test is held apart from selection.")
    add_heading(doc,"5 Internal Scaffold Test Evaluation")
    doc.add_paragraph("For each registered model, predictions were recomputed from the saved joblib artifact on saved scaffold-test rows. Bootstrap intervals quantify sampling variability of the test metrics, not per-compound uncertainty.")
    add_table(doc,["Gene","Model","n","R2","RMSE","MAE","Spearman"],[[r.gene,r.selected_model,str(int(r.n)),f"{r.R2:.3f}",f"{r.RMSE_pIC50:.3f}",f"{r.MAE_pIC50:.3f}",f"{r.Spearman_rho:.3f}"] for _,r in metrics.iterrows()])
    add_figure(doc,RESULTS/"figures"/"registered_model_scaffold_test_comparison.png","Figure 1. Recomputed internal scaffold-test performance across registered models.")
    add_figure(doc,RESULTS/"figures"/"activity_distribution_by_target.png","Figure 2. Curated pIC50 distributions by registered target.")
    add_heading(doc,"6 JAK Family Model Evaluation")
    doc.add_paragraph(f"The JAK-family model uses {int(selected.n)} saved scaffold-test compounds. Recomputed performance was R2 {selected.R2:.3f}, RMSE {selected.RMSE_pIC50:.3f}, MAE {selected.MAE_pIC50:.3f}, and Spearman rho {selected.Spearman_rho:.3f}. This is an internal performance estimate for a pooled pChEMBL endpoint.")
    add_table(doc,["Evaluation","R2","RMSE","MAE"],[[r['split'],f"{r.R2:.3f}",f"{r.RMSE_pIC50:.3f}",f"{r.MAE_pIC50:.3f}"] for _,r in jak_random.iterrows()])
    add_figure(doc,RESULTS/"figures"/"JAK_scaffold_test_diagnostics.png","Figure 3. JAK observed-versus-predicted, residual, and absolute-error diagnostics.")
    add_figure(doc,RESULTS/"figures"/"JAK_descriptor_distributions.png","Figure 4. RDKit descriptor distributions for the pooled JAK clean dataset.")
    add_figure(doc,RESULTS/"figures"/"JAK_descriptor_correlation_heatmap.png","Figure 5. Descriptor correlations for the pooled JAK clean dataset.")
    add_heading(doc,"7 Interpretability and Applicability Domain")
    doc.add_paragraph("A focused held-out permutation analysis evaluated the 20 fingerprint bits with highest native tree importance on a fixed 100-compound scaffold-test sample. This produces a local sensitivity diagnostic, not a causal molecular mechanism or a stability analysis across folds. The JAK applicability-domain analysis uses maximum Morgan Tanimoto similarity to the development set. The empirical 5th-percentile threshold is %.3f; it is an internal warning threshold, not a calibrated probability of correctness." % manifest["JAK_deep_dive"]["AD_test_similarity_threshold_5th_percentile"])
    add_figure(doc,RESULTS/"interpretability"/"JAK_targeted_permutation_importance.png","Figure 6. Targeted held-out permutation importance for the 20 highest native-importance fingerprint bits.")
    add_figure(doc,RESULTS/"applicability_domain"/"JAK_similarity_distribution.png","Figure 7. JAK scaffold-test nearest-neighbour similarity distribution.")
    add_heading(doc,"8 External Validation")
    doc.add_paragraph("An independent external validation requires a compatible endpoint and units, exact activity relation, assay context, provenance, structure and publication independence, sufficient sample size, and predefined performance criteria. Current external reports are retained as evidence inventory, but no registered model passed every gate.")
    add_table(doc,["Gene","Source","External n","RMSE","Spearman","Gate passed"],[[r.gene,str(r.source),str(r.n_nonoverlap),f"{r.RMSE_pIC50:.3f}" if pd.notna(r.RMSE_pIC50) else "-",f"{r.Spearman:.3f}" if pd.notna(r.Spearman) else "-",str(r.quality_gates_passed)] for _,r in external.iterrows()])
    add_heading(doc,"9 Reproducibility")
    doc.add_paragraph(f"The analysis used seed {manifest['random_seed']}, Python {manifest['python']}, RDKit {manifest['rdkit']}, scikit-learn {manifest['scikit_learn']}, NumPy {manifest['numpy']}, pandas {manifest['pandas']}, and SciPy {manifest['scipy']}. Exact model SHA256 values, recomputed predictions, residuals, diagnostics, audit tables, and validation inventory are stored under results/.")
    add_heading(doc,"10 Limitations and Conclusion")
    doc.add_paragraph("The models are useful for target-specific compound ranking within familiar chemistry, with explicit warnings for out-of-domain compounds. They do not provide clinical efficacy or safety predictions. Internal scaffold testing is stronger than a random split for chemical extrapolation but it is not a replacement for a context-matched independent assay set. The pooled JAK model cannot make isoform-specific claims. The correct current label is research screening, pending qualifying external validation and prospective experimental confirmation.")
    add_heading(doc,"References")
    for ref in ["OECD. Guidance Document on the Validation of (Quantitative) Structure-Activity Relationship [(Q)SAR] Models. OECD Series on Testing and Assessment No. 69, 2014. https://doi.org/10.1787/9789264085442-en", "RDKit documentation. The RDKit Book: Morgan and Feature Morgan fingerprints. https://www.rdkit.org/docs/RDKit_Book.html", "Rogers D, Hahn M. Extended-connectivity fingerprints. Journal of Chemical Information and Modeling. 2010;50:742-754. https://doi.org/10.1021/ci100050t"]: doc.add_paragraph(ref)
    doc.save(REPORTS/"QSAR_Research_Report.docx")
    print(REPORTS)

if __name__ == "__main__": main()
